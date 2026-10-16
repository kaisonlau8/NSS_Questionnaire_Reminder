"""车辆档案 Excel 导入（新版「客户车辆关系」导出，兼容旧交付档案格式）。

双文件导入见 parse_relation_file / parse_delivery_file / import_dual：
「客户车辆关系（DMS 导出）」提供归属门店，「车辆交付明细（数据门户导出）」提供交付时间，按 VIN 关联。
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from nss.jsonutil import load_json, save_json
from nss.paths import DATASOURCE_DIR, ROOT, UPLOAD_DIR, VEHICLES_PATH
from nss.settings import load_settings
from nss.time_utils import beijing_strftime, beijing_today

RELATION_FILE_LABEL = "客户车辆关系（DMS 导出）"
DELIVERY_FILE_LABEL = "车辆交付明细（数据门户导出）"
RELATION_SHEET_PREFIX = "客户车辆关系"
DELIVERY_SHEET_NAME = "预订订单数据明细"
RELATION_REQUIRED_HEADERS = ("VIN", "所属门店编码", "所属门店名称")
DELIVERY_REQUIRED_HEADERS = ("VIN码", "交付时间")

# 新版「客户车辆关系」导出：创建时间即交付日期；门店匹配只用「所属门店」列
NEW_FORMAT_MARKERS = ("创建时间", "所属门店名称")
NEW_REQUIRED_HEADERS = ("VIN", "车系名称", "创建时间", "所属门店名称")
# 旧版「车辆交付档案」
OLD_REQUIRED_HEADERS = ("VIN", "交付日期", "交付门店")
OLD_OPTIONAL_MAP = {
    "车牌号": "plate",
    "车系": "series",
    "车系名称": "series_name",
    "车型": "model",
    "售后车系": "aftersales_series",
    "数据类型": "data_type",
}


def parse_delivery_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    text = text.split()[0][:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def due_date_for(delivery: date, offset_days: int | None = None) -> date:
    days = offset_days if offset_days is not None else int(load_settings()["offset_days"])
    return delivery + timedelta(days=days)


def _header_map(headers: list[Any]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, raw in enumerate(headers):
        name = str(raw or "").strip()
        if name and name not in idx:
            idx[name] = i
    return idx


def _cell(row: tuple[Any, ...], idx: dict[str, int], header: str) -> str:
    if header not in idx:
        return ""
    return str(row[idx[header]] or "").strip()


def _sheet(wb):
    for ws in wb.worksheets:
        if str(ws.title).startswith("客户车辆关系"):
            return ws
    return wb.active


def parse_excel(path: Path) -> dict[str, Any]:
    """解析 Excel，返回 {"vehicles": [...], "stats": {total, imported, skipped_no_date, duplicates}}。

    新旧格式都只保留非个人信息字段；资格规则（车系/数据类型）不在导入时过滤，
    到期发送时再判定（见 nss/eligibility.py）。
    """
    wb = load_workbook(path, data_only=True)
    try:
        ws = _sheet(wb)
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise ValueError("Excel 为空")
    idx = _header_map(list(rows[0]))
    new_format = any(marker in idx for marker in NEW_FORMAT_MARKERS)
    required = NEW_REQUIRED_HEADERS if new_format else OLD_REQUIRED_HEADERS
    missing = [h for h in required if h not in idx]
    if missing:
        raise ValueError("Excel 缺少列: " + "、".join(missing) + f"。当前表头: {list(idx)}")
    date_header = "创建时间" if new_format else "交付日期"
    offset_days = int(load_settings()["offset_days"])
    vehicles: list[dict[str, Any]] = []
    stats = {"total": 0, "imported": 0, "skipped_no_date": 0, "duplicates": 0}
    seen: set[str] = set()
    for row in rows[1:]:
        if not row or not any(str(c or "").strip() for c in row):
            continue
        stats["total"] += 1
        vin = str(row[idx["VIN"]] or "").strip().upper()
        if not vin:
            continue
        delivery = parse_delivery_date(row[idx[date_header]])
        if not delivery:
            stats["skipped_no_date"] += 1
            continue
        if vin in seen:
            stats["duplicates"] += 1
            continue
        seen.add(vin)
        item: dict[str, Any] = {
            "vin": vin,
            "delivery_date": delivery.isoformat(),
            "due_date": due_date_for(delivery, offset_days).isoformat(),
            "store": "",
            "store_code": "",
            "delivery_store": "",
            "delivery_store_code": "",
            "plate": "",
            "series": "",
            "series_name": "",
            "model": "",
            "aftersales_series": "",
            "data_type": "",
        }
        if new_format:
            item.update(
                {
                    # 门店匹配只用所属门店，交付门店仅作展示参考
                    "store": _cell(row, idx, "所属门店名称"),
                    "store_code": _cell(row, idx, "所属门店编码"),
                    "delivery_store": _cell(row, idx, "交付门店名称"),
                    "delivery_store_code": _cell(row, idx, "交付门店编码"),
                    "plate": _cell(row, idx, "车牌号"),
                    "series_name": _cell(row, idx, "车系名称"),
                    "aftersales_series": _cell(row, idx, "售后车系"),
                    "model": _cell(row, idx, "车型名称"),
                    "data_type": _cell(row, idx, "数据类型"),
                }
            )
        else:
            item["store"] = _cell(row, idx, "交付门店")
            for header, key in OLD_OPTIONAL_MAP.items():
                if header in idx:
                    item[key] = _cell(row, idx, header)
        vehicles.append(item)
        stats["imported"] += 1
    return {"vehicles": vehicles, "stats": stats}


_CN_DATE_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")


def parse_portal_datetime(value: Any) -> date | None:
    """解析数据门户「交付时间」：2026年7月18日 10:28:54 / 2026-07-18 / 2026/07/18 / datetime，只取日期部分。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    m = _CN_DATE_RE.search(text)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    head = text.split()[0][:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(head, fmt).date()
        except ValueError:
            continue
    return None


def _delivery_sheet(wb):
    for ws in wb.worksheets:
        if str(ws.title) == DELIVERY_SHEET_NAME:
            return ws
    for ws in wb.worksheets:
        if "预订订单" in str(ws.title):
            return ws
    return wb.active


def _sheet_rows(path: Path, picker) -> list[list[Any]]:
    # 注意：这类导出在 openpyxl read_only 模式下 dimension 不可靠，统一用普通模式
    wb = load_workbook(path, data_only=True)
    try:
        ws = picker(wb)
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_relation_file(path: Path) -> dict[str, Any]:
    """解析「客户车辆关系（DMS 导出）」：只取归属/展示字段，绝不取交付日期。

    返回 {"records": {vin: fields}, "duplicates": n}；VIN 去重（首个生效）、大写去空格。
    """
    rows = _sheet_rows(path, _sheet)
    if not rows:
        raise ValueError(f"「{RELATION_FILE_LABEL}」Excel 为空")
    idx = _header_map(list(rows[0]))
    if "VIN码" in idx or "交付时间" in idx:
        raise ValueError(
            f"两个文件可能放反了：「{RELATION_FILE_LABEL}」里检测到 VIN码/交付时间 列，"
            f"这更像「{DELIVERY_FILE_LABEL}」，请交换后重试。当前表头: {list(idx)}"
        )
    missing = [h for h in RELATION_REQUIRED_HEADERS if h not in idx]
    if missing:
        raise ValueError(
            f"「{RELATION_FILE_LABEL}」缺少列: " + "、".join(missing) + f"。当前表头: {list(idx)}"
        )
    records: dict[str, dict[str, str]] = {}
    duplicates = 0
    for row in rows[1:]:
        if not row or not any(str(c or "").strip() for c in row):
            continue
        vin = str(row[idx["VIN"]] or "").strip().upper()
        if not vin:
            continue
        if vin in records:
            duplicates += 1
            continue
        records[vin] = {
            "vin": vin,
            "store": _cell(row, idx, "所属门店名称"),
            "store_code": _cell(row, idx, "所属门店编码"),
            "delivery_store": _cell(row, idx, "交付门店名称"),
            "delivery_store_code": _cell(row, idx, "交付门店编码"),
            "plate": _cell(row, idx, "车牌号"),
            "series": "",
            "series_name": _cell(row, idx, "车系名称"),
            "model": _cell(row, idx, "车型名称"),
            "aftersales_series": _cell(row, idx, "售后车系"),
            "data_type": _cell(row, idx, "数据类型"),
        }
    return {"records": records, "duplicates": duplicates}


def parse_delivery_file(path: Path) -> dict[str, Any]:
    """解析「车辆交付明细（数据门户导出）」：VIN码 + 交付时间，只取日期部分。

    VIN码或交付时间为空（未交付/已取消）的行忽略；VIN 去重（首个生效）。
    返回 {"dates": {vin: date}, "duplicates": n}。
    """
    rows = _sheet_rows(path, _delivery_sheet)
    if not rows:
        raise ValueError(f"「{DELIVERY_FILE_LABEL}」Excel 为空")
    idx = _header_map(list(rows[0]))
    if "所属门店编码" in idx or "所属门店名称" in idx:
        raise ValueError(
            f"两个文件可能放反了：「{DELIVERY_FILE_LABEL}」里检测到 所属门店编码/所属门店名称 列，"
            f"这更像「{RELATION_FILE_LABEL}」，请交换后重试。当前表头: {list(idx)}"
        )
    missing = [h for h in DELIVERY_REQUIRED_HEADERS if h not in idx]
    if missing:
        raise ValueError(
            f"「{DELIVERY_FILE_LABEL}」缺少列: " + "、".join(missing) + f"。当前表头: {list(idx)}"
        )
    dates: dict[str, date] = {}
    duplicates = 0
    for row in rows[1:]:
        if not row or not any(str(c or "").strip() for c in row):
            continue
        vin = str(row[idx["VIN码"]] or "").strip().upper()
        if not vin:
            continue
        day = parse_portal_datetime(row[idx["交付时间"]])
        if day is None:
            continue
        if vin in dates:
            duplicates += 1
            continue
        dates[vin] = day
    return {"dates": dates, "duplicates": duplicates}


def build_dual(relation_path: Path, delivery_path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """双文件按 VIN 关联：仅在两表都出现且有交付时间的车辆进入档案。"""
    relation = parse_relation_file(Path(relation_path))
    delivery = parse_delivery_file(Path(delivery_path))
    offset_days = int(load_settings()["offset_days"])
    records = relation["records"]
    dates = delivery["dates"]
    vehicles: list[dict[str, Any]] = []
    for vin, rec in records.items():
        day = dates.get(vin)
        if day is None:
            continue
        vehicles.append(
            {
                "vin": vin,
                "delivery_date": day.isoformat(),
                "due_date": due_date_for(day, offset_days).isoformat(),
                "store": rec["store"],
                "store_code": rec["store_code"],
                "delivery_store": rec["delivery_store"],
                "delivery_store_code": rec["delivery_store_code"],
                "plate": rec["plate"],
                "series": rec["series"],
                "series_name": rec["series_name"],
                "model": rec["model"],
                "aftersales_series": rec["aftersales_series"],
                "data_type": rec["data_type"],
            }
        )
    stats = {
        "relation_total": len(records),
        "delivery_total": len(dates),
        "matched": len(vehicles),
        "only_in_relation": len(records.keys() - dates.keys()),
        "only_in_delivery": len(dates.keys() - records.keys()),
        "duplicates": relation["duplicates"] + delivery["duplicates"],
        "unknown_store_codes": {"count": 0, "examples": []},
        "date_changed": {"count": 0, "examples": []},
        "past_due_unsent": {"count": 0, "examples": []},
    }
    return vehicles, stats


def _unknown_store_codes(vehicles: list[dict[str, Any]]) -> dict[str, Any]:
    from nss.bitable import load_stores  # 函数内导入，避免循环依赖

    known = {
        str(s.get("code") or "").strip().upper()
        for s in load_stores()
        if str(s.get("code") or "").strip()
    }
    seen: list[str] = []
    for v in vehicles:
        code = str(v.get("store_code") or "").strip().upper()
        if not code or code in known or code in seen:
            continue
        seen.append(code)
    return {"count": len(seen), "examples": seen[:20]}


def _date_changes(vehicles: list[dict[str, Any]], previous: list[dict[str, Any]]) -> dict[str, Any]:
    old_by_vin = {
        str(v.get("vin") or "").strip().upper(): str(v.get("delivery_date") or "")
        for v in previous
    }
    changes = [
        {"vin": v["vin"], "old": old, "new": v["delivery_date"]}
        for v in vehicles
        if (old := old_by_vin.get(v["vin"])) and old != v["delivery_date"]
    ]
    return {"count": len(changes), "examples": changes[:20]}


def _past_due_unsent(vehicles: list[dict[str, Any]]) -> dict[str, Any]:
    from nss.reminder import load_sent  # vehicles ↔ reminder 循环依赖，函数内导入

    today = beijing_today().isoformat()
    # 按 VIN 判断是否发过（交付日修正后 due_date 会变，旧 key 下发过也算已发）
    sent_vins = {str(i.get("key") or "").split(":")[0].upper() for i in load_sent().get("items") or []}
    items = []
    for v in vehicles:
        due = str(v.get("due_date") or "")
        if not due or due >= today:
            continue
        if v["vin"] in sent_vins:
            continue
        items.append({"vin": v["vin"], "due_date": due})
    return {"count": len(items), "examples": items[:20]}


def import_dual(
    relation_path: Path, delivery_path: Path, *, dry_run: bool = False
) -> dict[str, Any]:
    """双文件导入并覆盖 data/vehicles.json；dry_run 只统计不写盘。"""
    vehicles, stats = build_dual(relation_path, delivery_path)
    previous = load_vehicles()
    stats["unknown_store_codes"] = _unknown_store_codes(vehicles)
    stats["date_changed"] = _date_changes(vehicles, previous)
    stats["past_due_unsent"] = _past_due_unsent(vehicles)
    source = f"relation={Path(relation_path).name} | delivery={Path(delivery_path).name}"
    if dry_run:
        return {"dry_run": True, "source": source, "count": len(vehicles), "stats": stats}
    payload = save_vehicles(vehicles, source=source)
    payload.pop("vehicles", None)
    payload["stats"] = stats
    return payload


def save_vehicles(vehicles: list[dict[str, Any]], *, source: str) -> dict[str, Any]:
    payload = {
        "imported_at": beijing_strftime("%Y-%m-%d %H:%M:%S"),
        "source": source,
        "count": len(vehicles),
        "vehicles": vehicles,
    }
    save_json(VEHICLES_PATH, payload)
    return payload


def load_vehicles() -> list[dict[str, Any]]:
    data = load_json(VEHICLES_PATH, {})
    return data.get("vehicles") or []


def vehicles_meta() -> dict[str, Any]:
    data = load_json(VEHICLES_PATH, {})
    return {
        "imported_at": data.get("imported_at") or "",
        "source": data.get("source") or "",
        "count": int(data.get("count") or len(data.get("vehicles") or [])),
    }


def due_on(day: date | None = None) -> list[dict[str, Any]]:
    target = (day or beijing_today()).isoformat()
    return [v for v in load_vehicles() if v.get("due_date") == target]


def next_due_date() -> str:
    today = beijing_today().isoformat()
    dates = sorted({v.get("due_date") or "" for v in load_vehicles() if (v.get("due_date") or "") >= today})
    return dates[0] if dates else ""


def find_seed_excel() -> Path | None:
    candidates: list[Path] = []
    for folder in (UPLOAD_DIR, DATASOURCE_DIR, ROOT):
        if not folder.exists():
            continue
        candidates.extend(sorted(folder.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True))
    return candidates[0] if candidates else None


def import_excel(path: Path) -> dict[str, Any]:
    parsed = parse_excel(path)
    payload = save_vehicles(parsed["vehicles"], source=str(path))
    payload["stats"] = parsed["stats"]
    return payload

"""NSS 周报：纯函数统计、Excel、简报与默认 dry-run 的执行入口。"""

from __future__ import annotations

import argparse
import copy
import os
import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from nss.bitable import fetch_records, field_text
from nss.feishu_client import send_chat_text
from nss.jsonutil import load_json
from nss.matching import NO_REGION, enrich_vehicles
from nss.paths import DATA_DIR, WEEKLY_REPORT_DIR
from nss.settings import DEFAULT_SETTINGS, load_settings, normalize_weekly_report_settings
from nss.time_utils import BEIJING_TZ, beijing_today

REPORT_FILENAME_RE = re.compile(r"\d{4}-\d{2}-\d{2}_weekly\.xlsx")


def normalize_vin(value: Any) -> str:
    return re.sub(r"\s+", "", field_text(value)).upper()


def _datetime(value: Any) -> datetime | None:
    try:
        parsed = value if isinstance(value, datetime) else datetime.fromisoformat(field_text(value))
    except (ValueError, TypeError):
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=BEIJING_TZ)
    return parsed.astimezone(BEIJING_TZ)


def _monday(value: str | date) -> date:
    if isinstance(value, datetime):
        parsed = _datetime(value).date()
    elif isinstance(value, date):
        parsed = value
    else:
        parsed = date.fromisoformat(value)
    if parsed.weekday() != 0:
        raise ValueError("week_start / --week 必须是周一（YYYY-MM-DD）")
    return parsed


def previous_week_start(today: date | None = None) -> date:
    today = today or beijing_today()
    return today - timedelta(days=today.weekday() + 7)


def _items(payload: Any, key: str) -> list[dict[str, Any]]:
    items = payload.get(key, []) if isinstance(payload, dict) else payload
    if not isinstance(items, list) or any(not isinstance(item, dict) for item in items):
        raise ValueError(f"数据必须是对象列表或包含 {key} 列表的对象")
    return items


def _is_target(fields: dict[str, Any], filters: dict[str, Any]) -> bool:
    kind = field_text(fields.get(filters["type_field"]))
    period = field_text(fields.get(filters["period_field"]))
    if kind == filters["type_value"] and period == filters["period_value"]:
        return True
    combined = filters["combined_value"]
    # 单字段兼容不能盖过明确的「非 NPS / 非 90 天」结构化字段。
    return (
        kind in ("", filters["type_value"], combined)
        and period in ("", filters["period_value"], combined)
        and any(field_text(fields.get(key)) == combined for key in filters["combined_fields"])
    )


def _nps_score(fields: dict[str, Any], nps_field: str) -> int | str:
    keys = [nps_field] if nps_field else [
        key for key in fields
        if re.match(r"^\s*(?:第\s*)?1\s*(?:题|[.．、:：)）])", key)
    ]
    for key in keys:
        text = field_text(fields.get(key))
        try:
            score = float(text)
        except ValueError:
            continue
        if score.is_integer() and 0 <= score <= 10:
            return int(score)
    return ""


def _counts(items: list[dict[str, Any]]) -> dict[str, Any]:
    reminded = len(items)
    filled = sum(bool(item["filled"]) for item in items)
    return {"reminder_count": reminded, "filled_count": filled, "fill_rate": filled / reminded if reminded else 0.0}


def build_weekly_report(
    week_start: str | date,
    sent_log: dict[str, Any] | list[dict[str, Any]],
    bitable_rows: list[dict[str, Any]],
    *,
    filters: dict[str, Any] | None = None,
    nps_field: str = "",
) -> dict[str, Any]:
    """纯函数，不读配置/磁盘/网络，也不修改输入。

    提醒在本周按归一化 VIN 去重，保留最早提醒、合并接收人。填写按任务生成时间判定，
    下界是最早提醒当日 00:00，上界是下周一 00:00（不含）；提交时间仅展示。
    记录总数、未提醒记录数保留多维表逐条口径，另列未提醒的去重 VIN 台数。
    """
    monday = _monday(week_start)
    start = datetime.combine(monday, time.min, BEIJING_TZ)
    stop = start + timedelta(days=7)
    cfg = {**copy.deepcopy(DEFAULT_SETTINGS["weekly_report"]["filters"]), **(filters or {})}
    diagnostics = {"invalid_sent_rows": 0, "invalid_task_times": 0, "missing_record_vins": 0}
    sent = []
    for item in _items(sent_log, "items"):
        at = _datetime(item.get("sent_at"))
        vin = normalize_vin(item.get("vin"))
        if at is None or not vin:
            diagnostics["invalid_sent_rows"] += 1
        elif start <= at < stop:
            sent.append((at, vin, item))
    sent.sort(key=lambda entry: (entry[0], entry[1]))
    reminders: dict[str, dict[str, Any]] = {}
    for at, vin, item in sent:
        kept = reminders.setdefault(vin, {
            "vin": vin,
            "store": field_text(item.get("store")),
            "store_code": field_text(item.get("store_code")),
            "region": field_text(item.get("region")) or NO_REGION,
            "sent_at": at.isoformat(timespec="seconds"),
            "recipients": [],
        })
        seen = {p.get("open_id") or p.get("name") for p in kept["recipients"]}
        for person in item.get("recipients") or []:
            key = person.get("open_id") or person.get("name")
            if key and key not in seen:
                kept["recipients"].append(dict(person))
                seen.add(key)

    weekly_rows = []
    by_vin: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in _items(bitable_rows, "items"):
        fields = item.get("fields", item)
        if not isinstance(fields, dict):
            raise ValueError("多维表 fields 必须是对象")
        if not _is_target(fields, cfg):
            continue
        generated = _datetime(fields.get("任务生成时间"))
        if generated is None:
            diagnostics["invalid_task_times"] += 1
            continue
        if not start <= generated < stop:
            continue
        row = {
            "vin": normalize_vin(fields.get("VIN")),
            "store": field_text(fields.get("交付门店名称")),
            "region": field_text(fields.get("大区名称")) or NO_REGION,
            "task_generated_at": generated.strftime("%Y-%m-%d %H:%M:%S"),
            "submitted_at": field_text(fields.get("问卷提交时间")),
            "nps_score": _nps_score(fields, nps_field),
            "record_id": field_text(item.get("record_id") or item.get("id")),
        }
        if not row["vin"]:
            diagnostics["missing_record_vins"] += 1
        weekly_rows.append(row)
        by_vin[row["vin"]].append(row)
    weekly_rows.sort(key=lambda row: (row["task_generated_at"], row["vin"], row["record_id"]))
    for reminder in reminders.values():
        earliest = reminder["sent_at"][:10] + " 00:00:00"
        candidates = [row for row in by_vin[reminder["vin"]] if row["task_generated_at"] >= earliest]
        chosen = min(candidates, key=lambda row: (row["task_generated_at"], row["submitted_at"], row["record_id"])) if candidates else {}
        reminder.update({
            "filled": bool(chosen),
            "task_generated_at": chosen.get("task_generated_at", ""),
            "submitted_at": chosen.get("submitted_at", ""),
            "nps_score": chosen.get("nps_score", ""),
            "record_id": chosen.get("record_id", ""),
        })

    detail = list(reminders.values())
    unreminded = [row for row in weekly_rows if row["vin"] not in reminders]
    by_region: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_store: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in detail:
        by_region[item["region"]].append(item)
        by_store[(item["region"], item["store"] or "未知门店", item["store_code"])].append(item)
    diagnostics["duplicate_reminders"] = len(sent) - len(detail)
    diagnostics["no_region_count"] = sum(item["region"] == NO_REGION for item in detail)
    return {
        "week_start": monday.isoformat(),
        "week_end": (stop - timedelta(days=1)).date().isoformat(),
        "period_start": start.isoformat(),
        "period_end": (stop - timedelta(seconds=1)).isoformat(),
        **_counts(detail),
        "bitable_record_count": len(weekly_rows),
        "unreminded_count": len(unreminded),
        "unreminded_vin_count": len({row["vin"] for row in unreminded if row["vin"]}),
        "reminders": detail,
        "unreminded": unreminded,
        "summary": {
            "by_day": [
                {"date": day, **_counts([item for item in detail if item["sent_at"][:10] == day])}
                for day in ((monday + timedelta(days=i)).isoformat() for i in range(7))
            ],
            "by_region": [{"region": region, **_counts(items)} for region, items in sorted(by_region.items())],
            "by_store": [
                {"region": region, "store": store, "store_code": code, **_counts(items)}
                for (region, store, code), items in sorted(by_store.items())
            ],
        },
        "diagnostics": diagnostics,
    }


def render_brief(report: dict[str, Any], download_url: str) -> str:
    regions = [row for row in report["summary"]["by_region"] if row["region"] != NO_REGION]

    def ranking(reverse: bool) -> str:
        ranked = sorted(regions, key=lambda row: (
            -row["fill_rate"] if reverse else row["fill_rate"], -row["reminder_count"], row["region"],
        ))[:3]
        return "；".join(f"{row['region']} {row['filled_count']}/{row['reminder_count']}（{row['fill_rate']:.1%}）" for row in ranked) or "暂无可归属大区"

    lines = [
        f"NSS 周报｜{report['week_start']} ~ {report['week_end']}（北京时间）",
        f"提醒 {report['reminder_count']} 台（按 VIN 去重），已填写 {report['filled_count']} 台（{report['fill_rate']:.1%}）",
        f"大区 Top：{ranking(True)}",
        f"大区 Bottom：{ranking(False)}",
        f"本周 NPS-90 共 {report['bitable_record_count']} 条；未提醒但已填写 {report['unreminded_count']} 条（{report['unreminded_vin_count']} 台）",
    ]
    if report["diagnostics"]["no_region_count"]:
        lines.append(f"无区域：{report['diagnostics']['no_region_count']} 台（未纳入大区排名）")
    lines.append(f"明细下载（需登录）：{download_url}")
    return "\n".join(lines)


def _append_row(sheet, values: list[Any]) -> None:
    sheet.append(values)
    # 源数据为文本：保留原文，不把 = 开头的门店/接收人等解释成 Excel 公式。
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1677FF")
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_weekly_excel(report: dict[str, Any], path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    detail = workbook.active
    detail.title = "提醒明细"
    _append_row(detail, ["VIN", "门店", "大区", "提醒时间", "接收人", "是否填写", "任务生成时间", "问卷提交时间", "NPS 分(第1题)"])
    for item in report["reminders"]:
        people = "、".join(
            f"{p.get('name') or p.get('open_id') or ''}（{p.get('role') or '接收人'}）"
            for p in item["recipients"]
        )
        _append_row(detail, [
            item["vin"], item["store"], item["region"], item["sent_at"], people,
            "是" if item["filled"] else "否", item["task_generated_at"], item["submitted_at"], item["nps_score"],
        ])
    _style_sheet(detail, [23, 34, 15, 29, 42, 12, 23, 23, 18])

    unreminded = workbook.create_sheet("未提醒但已填写")
    _append_row(unreminded, ["VIN", "门店", "大区", "任务生成时间", "问卷提交时间", "NPS 分(第1题)", "记录ID"])
    for item in report["unreminded"]:
        _append_row(unreminded, [item[key] for key in (
            "vin", "store", "region", "task_generated_at", "submitted_at", "nps_score", "record_id",
        )])
    _style_sheet(unreminded, [23, 34, 15, 23, 23, 18, 24])

    summary = workbook.create_sheet("汇总")
    _append_row(summary, ["维度", "日期 / 大区 / 门店", "大区", "提醒数", "填写数", "填写率"])
    _append_row(summary, ["全周", f"{report['week_start']} ~ {report['week_end']}", "", report["reminder_count"], report["filled_count"], report["fill_rate"]])
    for dimension, key, label in (("by_day", "date", "按天"), ("by_region", "region", "按大区"), ("by_store", "store", "按门店")):
        for item in report["summary"][dimension]:
            name = item[key]
            if dimension == "by_store" and item["store_code"]:
                name += f"（{item['store_code']}）"
            _append_row(summary, [label, name, item.get("region", ""), item["reminder_count"], item["filled_count"], item["fill_rate"]])
    _style_sheet(summary, [28, 56, 18, 14, 14, 16])
    for row in summary.iter_rows(min_row=2):
        row[5].number_format = "0.0%"
    for label, value in (
        ("周期（Asia/Shanghai）", f"{report['period_start']} ~ {report['period_end']}"),
        ("本周 NPS-90 记录数", report["bitable_record_count"]),
        ("未提醒但已填写记录数", report["unreminded_count"]),
        ("未提醒但已填写 VIN 数", report["unreminded_vin_count"]),
        ("重复提醒合并条数", report["diagnostics"]["duplicate_reminders"]),
        ("本周记录缺失 VIN 条数", report["diagnostics"]["missing_record_vins"]),
        ("无法解析的提醒条数（全日志）", report["diagnostics"]["invalid_sent_rows"]),
        ("NPS-90 无效任务时间（全数据）", report["diagnostics"]["invalid_task_times"]),
        ("无区域提醒台数", report["diagnostics"]["no_region_count"]),
        ("填写口径", "任务生成时间在最早提醒当日 00:00 至本周日内；提交时间仅展示；同 VIN 取最早合格记录"),
        ("统计口径", "提醒/填写按 VIN 去重，按天归属最早提醒日；未提醒记录指 VIN 不在本周提醒集合，空 VIN 也计入记录数"),
        ("归属口径", "提醒门店按现有 DMS 店代码匹配，大区只认 DMS reportdb；未提醒明细门店/大区为问卷源数据"),
    ):
        _append_row(summary, [label, value])
    temporary = path.with_suffix(path.suffix + ".tmp")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return path


def _enrich_sent_log(sent_log: Any, data_dir: Path, region_db: str | Path | None) -> dict[str, Any]:
    vehicles = _items(load_json(data_dir / "vehicles.json", {}), "vehicles")
    stores = _items(load_json(data_dir / "stores.json", {}), "stores")
    by_vin = {
        normalize_vin(vehicle.get("vin")): vehicle
        for vehicle in enrich_vehicles(vehicles, stores, region_db=region_db)
    }
    items = []
    for item in _items(sent_log, "items"):
        vehicle = by_vin.get(normalize_vin(item.get("vin")))
        context = {} if vehicle is None else {
            "store": vehicle.get("store") or vehicle.get("matched_store") or "",
            "store_code": vehicle.get("store_code") or "",
            "region": vehicle["region"],
        }
        items.append({**item, **context})
    return {"items": items}


def run_weekly_report(
    week_start: str | date | None = None,
    *,
    dry_run: bool | None = None,
    sent_log_path: Path | str | None = None,
    rows_json_path: Path | str | None = None,
    data_dir: Path | str | None = None,
    out_dir: Path | str | None = None,
    region_db: Path | str | None = None,
) -> dict[str, Any]:
    settings = load_settings()
    cfg = normalize_weekly_report_settings(settings.get("weekly_report", {}))
    dry_run = cfg["dry_run"] if dry_run is None else dry_run
    if not isinstance(dry_run, bool):
        raise ValueError("dry_run 必须是布尔值")
    monday = _monday(week_start if week_start is not None else previous_week_start())
    source_dir = Path(data_dir or os.getenv("NSS_WEEKLY_DATA_DIR") or (
        Path(sent_log_path).parent if sent_log_path is not None else DATA_DIR
    )).expanduser()
    sent_path = Path(sent_log_path) if sent_log_path is not None else source_dir / "sent_log.json"
    if not sent_path.is_file():
        raise FileNotFoundError(f"找不到提醒日志：{sent_path}")
    sent_log = _enrich_sent_log(load_json(sent_path, {}), source_dir, region_db)
    if rows_json_path is not None:
        rows_path = Path(rows_json_path)
        if not rows_path.is_file():
            raise FileNotFoundError(f"找不到问卷数据：{rows_path}")
        rows = _items(load_json(rows_path, []), "items")
    else:
        source = cfg["source"]
        rows = fetch_records(source["table_id"], source["view_id"], app_token=source["app_token"], raw_items=True)
    report = build_weekly_report(monday, sent_log, rows, filters=cfg["filters"], nps_field=cfg["nps_field"])
    directory = Path(out_dir) if out_dir is not None else WEEKLY_REPORT_DIR
    filename = f"{monday.isoformat()}_weekly.xlsx"
    excel_path = write_weekly_excel(report, directory / filename)
    base = (cfg["link_base"] or settings.get("public_base_url") or "").rstrip("/")
    download_url = f"{base}/weekly-reports/{quote(filename)}"
    brief = render_brief(report, download_url)
    brief_path = directory / f"{monday.isoformat()}_brief.txt"
    temporary = brief_path.with_suffix(".txt.tmp")
    temporary.write_text(brief + "\n", encoding="utf-8")
    temporary.replace(brief_path)
    # 只在明确关闭 dry-run 时进入发送分支；离线 dry-run 不触发任何 API。
    message = None
    if not dry_run:
        message = send_chat_text(cfg["chat_id"], brief)
    return {
        "ok": True, "dry_run": dry_run, "report": report, "brief": brief,
        "excel_path": str(excel_path), "brief_path": str(brief_path),
        "download_url": download_url, "message": message,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="NSS 上一完整自然周周报（默认 dry-run）")
    parser.add_argument("--week", help="周一日期 YYYY-MM-DD；省略为上一完整周")
    parser.add_argument("--dry-run", action="store_true", default=None, help="强制只生成本地文件，绝不发送消息")
    parser.add_argument("--sent-log", type=Path, help="本地 sent_log.json")
    parser.add_argument("--rows-json", type=Path, help="本地 records API items 列表（或 fields 字典列表）")
    parser.add_argument("--data-dir", type=Path, help="只读车辆/门店快照目录；默认日志所在目录或项目 data/")
    parser.add_argument("--region-db", type=Path, help="只读 DMS reportdb 快照；默认 NSS_DMS_REGION_DB")
    parser.add_argument("--out", type=Path, help="产物目录；默认 data/weekly_reports/")
    args = parser.parse_args(argv)
    try:
        result = run_weekly_report(
            args.week, dry_run=args.dry_run, sent_log_path=args.sent_log,
            rows_json_path=args.rows_json, data_dir=args.data_dir,
            out_dir=args.out, region_db=args.region_db,
        )
    except (ValueError, OSError, RuntimeError) as exc:
        parser.exit(1, f"周报生成失败：{exc}\n")
    print(f"模式：{'dry-run' if result['dry_run'] else '已发送'}")
    print(result["brief"])
    print(f"Excel：{result['excel_path']}")
    print(f"简报：{result['brief_path']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

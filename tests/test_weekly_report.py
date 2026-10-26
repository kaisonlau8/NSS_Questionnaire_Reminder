"""周报口径、只读分页、Excel、dry-run、调度及真实 SSO 请求钩子的回归测试。"""

import copy
import json
import sqlite3
from datetime import date, datetime
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

import app as app_module
import nss.bitable as bitable
import nss.feishu_client as feishu
import nss.scheduler as scheduler
import nss.settings as settings_mod
import nss.weekly_report as weekly
from mhero_auth.core import AuthConfig
from nss.time_utils import BEIJING_TZ

WEEK = "2026-10-19"
QUESTION = "1.综合考虑您的用车感受，您有多大可能向您的朋友或同事推荐这款车？"


def sent(vin, at="2026-10-19T09:00:00+08:00", **extra):
    return {"vin": vin, "sent_at": at, "due_date": "1999-01-01", "recipients": [], **extra}


def row(vin, generated="2026-10-20 10:00:00", *, record_id="rec-test", **extra):
    return {"record_id": record_id, "fields": {
        "VIN": vin, "问卷类型": "NPS调研", "调研时间": "90天",
        "任务生成时间": generated, "问卷提交时间": "2026-10-20 11:00:00", QUESTION: "9",
        **extra,
    }}


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(settings_mod, "SETTINGS_PATH", tmp_path / "settings.json")
    monkeypatch.setattr(settings_mod, "SETTINGS_EXAMPLE_PATH", tmp_path / "settings.example.json")
    for name in (
        "NSS_WEEKLY_BITABLE_APP_TOKEN", "NSS_WEEKLY_BITABLE_TABLE_ID", "NSS_WEEKLY_BITABLE_VIEW_ID",
        "NSS_WEEKLY_CHAT_ID", "NSS_WEEKLY_LINK_BASE", "NSS_PUBLIC_BASE_URL", "NSS_WEEKLY_DATA_DIR",
    ):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("NSS_DMS_REGION_DB", str(tmp_path / "missing.sqlite3"))


def test_week_boundaries_use_beijing_sent_at_and_task_time():
    logs = [
        sent("before", "2026-10-18T23:59:59+08:00"),
        sent("start", "2026-10-18T16:00:00Z"),  # 北京周一 00:00
        sent("end", "2026-10-25T23:59:59+08:00"),
        sent("after", "2026-10-26T00:00:00+08:00"),
    ]
    rows = [
        row("start", "2026-10-19 00:00:00"),
        row("end", "2026-10-25 23:59:59"),
        row("other-before", "2026-10-18 23:59:59"),
        row("other-after", "2026-10-26 00:00:00"),
    ]
    result = weekly.build_weekly_report(WEEK, {"items": logs}, rows)
    assert result["reminder_count"] == result["filled_count"] == result["bitable_record_count"] == 2
    assert result["fill_rate"] == 1
    assert result["period_start"] == "2026-10-19T00:00:00+08:00"
    assert result["period_end"] == "2026-10-25T23:59:59+08:00"
    assert result["summary"]["by_day"][0]["reminder_count"] == 1
    assert result["summary"]["by_day"][-1]["filled_count"] == 1


def test_vin_normalization_dedup_earliest_reminder_and_recipient_union():
    people = [{"name": "经理", "open_id": "ou_m", "role": "manager"}]
    logs = [
        sent(" ab c ", "2026-10-21T09:00:00+08:00", recipients=people),
        sent("A\tBC", "2026-10-20T09:00:00+08:00", recipients=[
            {"name": "督导", "open_id": "ou_s", "role": "supervisor"}, *people,
        ]),
    ]
    result = weekly.build_weekly_report(WEEK, logs, [row("a b\nc", "2026-10-20 00:00:00")])
    assert result["reminder_count"] == result["filled_count"] == 1
    assert result["diagnostics"]["duplicate_reminders"] == 1
    item = result["reminders"][0]
    assert item["vin"] == "ABC"
    assert item["sent_at"] == "2026-10-20T09:00:00+08:00"
    assert [p["open_id"] for p in item["recipients"]] == ["ou_s", "ou_m"]


def test_tasks_before_sent_date_do_not_count_but_same_date_before_send_does():
    result = weekly.build_weekly_report(WEEK, [
        sent("before", "2026-10-22T09:00:00+08:00"),
        sent("same-day", "2026-10-22T09:00:00+08:00"),
        sent("next-week"),
    ], [
        row("before", "2026-10-21 23:59:59"),
        row("same-day", "2026-10-22 00:00:00", **{"问卷提交时间": ""}),
        row("next-week", "2026-10-26 00:00:00"),
    ])
    filled = {item["vin"]: item["filled"] for item in result["reminders"]}
    assert filled == {"NEXT-WEEK": False, "BEFORE": False, "SAME-DAY": True}
    assert result["filled_count"] == 1
    assert result["fill_rate"] == pytest.approx(1 / 3)
    assert result["bitable_record_count"] == 2
    assert result["unreminded_count"] == 0  # 同周提醒 VIN 的早期任务也不是「未提醒」。


def test_target_filter_excludes_other_surveys_and_supports_combined_single_field():
    rows = [
        row("split"),
        row("combined-type", **{"问卷类型": "NPS调研-90天", "调研时间": ""}),
        row("combined-name", **{"问卷类型": "", "调研时间": "", "问卷名称": "NPS调研-90天"}),
        row("other-type", **{"问卷类型": "CSI调研", "问卷名称": "NPS调研-90天"}),
        row("other-period", **{"调研时间": "30天", "问卷名称": "NPS调研-90天"}),
        row("bad-time", "无日期"),
    ]
    result = weekly.build_weekly_report(WEEK, [], rows)
    assert result["bitable_record_count"] == 3
    assert {item["vin"] for item in result["unreminded"]} == {"SPLIT", "COMBINED-TYPE", "COMBINED-NAME"}
    assert result["diagnostics"]["invalid_task_times"] == 1


def test_configurable_filter_fields_values_and_nps_question():
    filters = {
        "type_field": "类型", "type_value": "售后", "period_field": "周期", "period_value": "100天",
        "combined_value": "售后-100天", "combined_fields": ["名称"],
    }
    rows = [{"VIN": "v1", "类型": "售后", "周期": "100天", "任务生成时间": "2026-10-20 10:00:00", "自定义推荐分": "0"}]
    result = weekly.build_weekly_report(WEEK, [sent("v1")], rows, filters=filters, nps_field="自定义推荐分")
    assert result["filled_count"] == 1
    assert result["reminders"][0]["nps_score"] == 0


def test_unreminded_preserves_record_counts_and_separate_unique_vin_counts():
    rows = [
        row("reminded"), row("outside-cohort", record_id="rec-one"),
        row(" outside-cohort ", record_id="rec-two"), row(None, record_id="rec-no-vin"),
    ]
    logs = [sent("reminded"), sent("outside-cohort", "2026-10-12T09:00:00+08:00")]
    result = weekly.build_weekly_report(WEEK, logs, rows)
    assert result["bitable_record_count"] == 4
    assert result["unreminded_count"] == 3
    assert result["unreminded_vin_count"] == 1
    assert result["diagnostics"]["missing_record_vins"] == 1
    assert {item["record_id"] for item in result["unreminded"]} == {"rec-one", "rec-two", "rec-no-vin"}


def test_one_vin_multiple_tasks_chooses_first_qualifying_record_not_first_input():
    rows = [
        row("v1", "2026-10-24 10:00:00", record_id="late", **{QUESTION: "10"}),
        row("v1", "2026-10-20 10:00:00", record_id="early", **{QUESTION: "0", "10.其他题": "8"}),
    ]
    result = weekly.build_weekly_report(WEEK, [sent("v1")], rows)
    assert result["reminder_count"] == result["filled_count"] == 1
    assert result["bitable_record_count"] == 2
    assert result["reminders"][0]["record_id"] == "early"
    assert result["reminders"][0]["nps_score"] == 0


def test_pure_function_does_not_mutate_inputs_or_load_settings(monkeypatch):
    inputs = ({"items": [sent("a", store="店A", region="大区A")]}, [row(" a ")])
    before = copy.deepcopy(inputs)
    monkeypatch.setattr(weekly, "load_settings", Mock(side_effect=AssertionError("pure")))
    monkeypatch.setattr(weekly, "fetch_records", Mock(side_effect=AssertionError("pure")))
    report = weekly.build_weekly_report(WEEK, *inputs)
    assert inputs == before
    assert json.loads(json.dumps(report))["filled_count"] == 1


def test_summary_daily_region_and_store_use_same_deduplicated_cohort():
    logs = [
        sent("a", store="店A", store_code="S1", region="北区"),
        sent("b", store="店B", store_code="S2", region="北区"),
        sent("c", "2026-10-20T09:00:00+08:00", store="店A", store_code="S3", region="南区"),
        sent("a", "2026-10-21T09:00:00+08:00", store="店A", store_code="S1", region="北区"),
    ]
    report = weekly.build_weekly_report(WEEK, logs, [row("a"), row("c")])
    for dimension in report["summary"].values():
        assert sum(item["reminder_count"] for item in dimension) == 3
        assert sum(item["filled_count"] for item in dimension) == 2
    regions = {item["region"]: item for item in report["summary"]["by_region"]}
    assert regions["北区"]["fill_rate"] == 0.5
    assert regions["南区"]["fill_rate"] == 1
    assert len(report["summary"]["by_store"]) == 3
    brief = weekly.render_brief(report, "/weekly-reports/example.xlsx")
    assert "大区 Top：南区 1/1（100.0%）" in brief
    assert "大区 Bottom：北区 1/2（50.0%）" in brief
    assert "提醒 3 台（按 VIN 去重），已填写 2 台（66.7%）" in brief


def test_empty_week_and_invalid_dates_are_explicit():
    report = weekly.build_weekly_report(WEEK, [sent("bad", "invalid"), sent("")], [])
    assert report["fill_rate"] == 0
    assert report["diagnostics"]["invalid_sent_rows"] == 2
    assert len(report["summary"]["by_day"]) == 7
    assert "暂无可归属大区" in weekly.render_brief(report, "/weekly-reports")
    with pytest.raises(ValueError, match="周一"):
        weekly.build_weekly_report("2026-10-20", [], [])


def test_excel_contains_required_details_raw_records_and_numeric_rates(tmp_path):
    report = weekly.build_weekly_report(WEEK, [sent(
        "a", store="=SUM(1,2)", region="北区", recipients=[{"name": "经理A", "role": "manager"}],
    )], [row("a", **{QUESTION: "0"}), row("b"), row("b", record_id="second")])
    path = weekly.write_weekly_excel(report, tmp_path / "report.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == ["提醒明细", "未提醒但已填写", "汇总"]
        detail = workbook["提醒明细"]
        assert detail.max_row == 2
        assert detail["A2"].value == "A"
        assert detail["B2"].value == "=SUM(1,2)" and detail["B2"].data_type == "s"
        assert detail["E2"].value == "经理A（manager）"
        assert detail["F2"].value == "是" and detail["I2"].value == 0
        assert workbook["未提醒但已填写"].max_row == 3
        assert workbook["汇总"]["F2"].value == 1
        assert workbook["汇总"]["F2"].number_format == "0.0%"
        assert detail.freeze_panes == "A2"
    finally:
        workbook.close()


@pytest.mark.parametrize("hour,minute", [
    (0, 0), (0, 5), (23, 55), (23, 59),
    (8, 25), (8, 30), (8, 35), (8, 55), (9, 0), (9, 5),
    (11, 55), (12, 0), (12, 5),
])
def test_schedule_rejects_reserved_windows_including_midnight_wrap(hour, minute):
    with pytest.raises(ValueError, match="前后 5 分钟"):
        settings_mod.validate_weekly_report_schedule({"weekday": 0, "hour": hour, "minute": minute})


@pytest.mark.parametrize("hour,minute", [(0, 6), (23, 54), (8, 24), (8, 36), (8, 54), (9, 6), (9, 30), (11, 54), (12, 6)])
def test_schedule_accepts_times_outside_reserved_windows(hour, minute):
    settings_mod.validate_weekly_report_schedule({"weekday": 0, "hour": hour, "minute": minute})


@pytest.mark.parametrize("config", [
    {"weekday": -1}, {"weekday": 7}, {"hour": 24}, {"minute": 60},
    {"hour": 9.5}, {"weekday": True}, {"minute": "x"},
])
def test_schedule_rejects_invalid_values(config):
    with pytest.raises(ValueError):
        settings_mod.validate_weekly_report_schedule(config)


def test_settings_defaults_merge_nested_updates_and_reject_before_saving():
    defaults = settings_mod.load_settings()["weekly_report"]
    assert defaults["enabled"] is False and defaults["dry_run"] is True
    assert (defaults["weekday"], defaults["hour"], defaults["minute"]) == (0, 9, 30)
    settings_mod.save_settings({"weekly_report": {"filters": {"period_value": "100天"}}})
    settings_mod.save_settings({"weekly_report": {"minute": 31}})
    current = settings_mod.load_settings()["weekly_report"]
    assert current["filters"]["period_value"] == "100天"
    assert current["source"] == defaults["source"]
    before = settings_mod.SETTINGS_PATH.read_bytes()
    with pytest.raises(ValueError):
        settings_mod.save_settings({"weekly_report": {"hour": 9, "minute": 0}})
    assert settings_mod.SETTINGS_PATH.read_bytes() == before
    with pytest.raises(ValueError, match="true 或 false"):
        settings_mod.save_settings({"weekly_report": {"dry_run": "false"}})


def test_source_chat_and_link_settings_support_environment_overrides(monkeypatch):
    monkeypatch.setenv("NSS_WEEKLY_BITABLE_APP_TOKEN", "local-app")
    monkeypatch.setenv("NSS_WEEKLY_BITABLE_TABLE_ID", "local-table")
    monkeypatch.setenv("NSS_WEEKLY_BITABLE_VIEW_ID", "")
    monkeypatch.setenv("NSS_WEEKLY_CHAT_ID", "local-chat")
    monkeypatch.setenv("NSS_WEEKLY_LINK_BASE", "http://127.0.0.1:9005")
    cfg = settings_mod.load_settings()["weekly_report"]
    assert cfg["source"] == {"app_token": "local-app", "table_id": "local-table", "view_id": ""}
    assert cfg["chat_id"] == "local-chat"
    assert cfg["link_base"] == "http://127.0.0.1:9005"
    assert cfg["dry_run"] is True and cfg["enabled"] is False


def test_scheduler_fires_previous_completed_week_once_and_defaults_to_disabled(monkeypatch):
    run = Mock()
    monkeypatch.setattr(scheduler, "run_weekly_report", run)
    settings = settings_mod.load_settings()
    now = datetime(2026, 10, 26, 9, 30, tzinfo=BEIJING_TZ)
    fired = {}
    scheduler._run_weekly_if_due(now, settings, fired)
    run.assert_not_called()
    settings["weekly_report"]["enabled"] = True
    scheduler._run_weekly_if_due(now.replace(minute=29), settings, fired)
    scheduler._run_weekly_if_due(now.replace(day=27), settings, fired)
    run.assert_not_called()
    monkeypatch.setitem(scheduler._state, "last_weekly_report", "")
    scheduler._run_weekly_if_due(now, settings, fired)
    scheduler._run_weekly_if_due(now.replace(second=10), settings, fired)
    run.assert_called_once_with(date(2026, 10, 19), dry_run=True)
    assert scheduler._state["last_weekly_report"]


def test_scheduler_rejects_bad_schedule_before_trigger(monkeypatch):
    cfg = settings_mod.load_settings()
    cfg["weekly_report"].update(hour=9, minute=0)
    monkeypatch.setattr(scheduler, "load_settings", lambda: cfg)
    with pytest.raises(ValueError):
        scheduler._run_weekly_if_due(datetime(2026, 10, 26, 9, 0), cfg, {})


def test_scheduler_start_rejects_invalid_file_before_thread_creation(monkeypatch):
    settings_mod.SETTINGS_PATH.write_text(json.dumps({"weekly_report": {"hour": 9, "minute": 0}}))
    thread = Mock()
    monkeypatch.setattr(scheduler.threading, "Thread", thread)
    monkeypatch.setitem(scheduler._state, "running", False)
    with pytest.raises(ValueError, match="前后 5 分钟"):
        scheduler.start_scheduler()
    thread.assert_not_called()
    assert scheduler._state["running"] is False


def offline_inputs(tmp_path):
    source = tmp_path / "snapshot"
    source.mkdir()
    (source / "sent_log.json").write_text(json.dumps({"items": [sent("a")]}), encoding="utf-8")
    (source / "vehicles.json").write_text(json.dumps({"vehicles": [{"vin": "a", "store": "DMS门店", "store_code": "S001"}]}), encoding="utf-8")
    (source / "stores.json").write_text(json.dumps({"stores": [{"name": "门店清单", "code": "S001", "region": "飞书区域不能覆盖DMS"}]}), encoding="utf-8")
    rows_path = tmp_path / "rows.json"
    rows_path.write_text(json.dumps([row("a")]), encoding="utf-8")
    return source, rows_path


@pytest.mark.parametrize("force_dry_run,configured_dry_run", [(None, True), (True, False)])
def test_dry_run_writes_brief_and_excel_without_any_send_or_auth_call(tmp_path, monkeypatch, force_dry_run, configured_dry_run):
    source, rows_path = offline_inputs(tmp_path)
    settings_mod.save_settings({"weekly_report": {"dry_run": configured_dry_run}})
    send = Mock(side_effect=AssertionError("dry-run must not send"))
    post = Mock(side_effect=AssertionError("dry-run must not call auth / send API"))
    fetch = Mock(side_effect=AssertionError("offline must not fetch"))
    monkeypatch.setattr(weekly, "send_chat_text", send)
    monkeypatch.setattr(weekly, "fetch_records", fetch)
    monkeypatch.setattr(feishu.requests, "post", post)
    result = weekly.run_weekly_report(
        WEEK, dry_run=force_dry_run, sent_log_path=source / "sent_log.json",
        rows_json_path=rows_path, out_dir=tmp_path / "out",
    )
    assert result["dry_run"] is True and result["message"] is None
    assert result["report"]["reminder_count"] == result["report"]["filled_count"] == 1
    assert result["report"]["reminders"][0]["store"] == "DMS门店"
    assert result["report"]["reminders"][0]["region"] == "无区域"
    assert (tmp_path / "out" / f"{WEEK}_brief.txt").read_text() == result["brief"] + "\n"
    assert (tmp_path / "out" / f"{WEEK}_weekly.xlsx").is_file()
    assert not (source / "weekly_reports").exists()
    send.assert_not_called()
    post.assert_not_called()
    fetch.assert_not_called()


def test_runner_uses_explicit_dms_region_snapshot_readonly(tmp_path):
    source, rows_path = offline_inputs(tmp_path)
    region_db = tmp_path / "reportdb.sqlite3"
    with sqlite3.connect(region_db) as conn:
        conn.execute("CREATE TABLE stores (code TEXT, region TEXT)")
        conn.execute("INSERT INTO stores VALUES ('s001', 'DMS北区')")
    before = region_db.read_bytes()
    result = weekly.run_weekly_report(
        WEEK, dry_run=True, data_dir=source, rows_json_path=rows_path,
        region_db=region_db, out_dir=tmp_path / "out",
    )
    assert result["report"]["reminders"][0]["region"] == "DMS北区"
    assert "大区 Top：DMS北区 1/1（100.0%）" in result["brief"]
    assert region_db.read_bytes() == before


def test_cli_force_dry_run_and_public_link_fallback(tmp_path, monkeypatch, capsys):
    source, rows_path = offline_inputs(tmp_path)
    settings_mod.save_settings({"weekly_report": {"dry_run": False}, "public_base_url": "http://127.0.0.1:9005/"})
    send = Mock()
    monkeypatch.setattr(weekly, "send_chat_text", send)
    assert weekly.main([
        "--week", WEEK, "--dry-run", "--sent-log", str(source / "sent_log.json"),
        "--rows-json", str(rows_path), "--out", str(tmp_path / "out"),
    ]) == 0
    assert f"http://127.0.0.1:9005/weekly-reports/{WEEK}_weekly.xlsx" in capsys.readouterr().out
    send.assert_not_called()


def test_runner_uses_readonly_source_and_explicit_group_send_is_mocked(tmp_path, monkeypatch):
    source, _ = offline_inputs(tmp_path)
    settings_mod.save_settings({"weekly_report": {
        "dry_run": False, "chat_id": "oc-local-test", "link_base": "http://127.0.0.1:9005",
    }})
    monkeypatch.setenv("NSS_WEEKLY_DATA_DIR", str(source))
    monkeypatch.setattr(weekly, "WEEKLY_REPORT_DIR", tmp_path / "reports")
    fetch = Mock(return_value=[row("a")])
    send = Mock(return_value={"message_id": "mock-message"})
    monkeypatch.setattr(weekly, "fetch_records", fetch)
    monkeypatch.setattr(weekly, "send_chat_text", send)
    result = weekly.run_weekly_report(WEEK)
    cfg = settings_mod.load_settings()["weekly_report"]["source"]
    fetch.assert_called_once_with(cfg["table_id"], cfg["view_id"], app_token=cfg["app_token"], raw_items=True)
    send.assert_called_once_with("oc-local-test", result["brief"])
    assert result["message"] == {"message_id": "mock-message"}


def test_missing_local_input_fails_without_network(tmp_path, monkeypatch):
    fetch = Mock()
    monkeypatch.setattr(weekly, "fetch_records", fetch)
    with pytest.raises(FileNotFoundError):
        weekly.run_weekly_report(WEEK, sent_log_path=tmp_path / "missing.json", rows_json_path=tmp_path / "rows.json")
    fetch.assert_not_called()


def response(payload):
    return Mock(json=Mock(return_value=payload))


def test_bitable_records_are_get_only_page_size_500_and_preserve_ids(monkeypatch):
    get = Mock(side_effect=[
        response({"code": 0, "data": {"items": [row("a")], "has_more": True, "page_token": "page2"}}),
        response({"code": 0, "data": {"items": [row("b")], "has_more": False}}),
    ])
    post = Mock(side_effect=AssertionError("readonly"))
    monkeypatch.setattr(bitable.requests, "get", get)
    monkeypatch.setattr(bitable.requests, "post", post)
    monkeypatch.setattr(bitable, "auth_headers", lambda: {"Authorization": "Bearer mock-token"})
    records = bitable.fetch_records("tbl-test", "view-test", app_token="app-test", raw_items=True)
    assert len(records) == 2 and records[0]["record_id"] == "rec-test"
    for call in get.call_args_list:
        assert call.args[0].endswith("/apps/app-test/tables/tbl-test/records")
        assert call.kwargs["params"]["page_size"] == 500
        assert call.kwargs["params"]["view_id"] == "view-test"
    assert "page_token" not in get.call_args_list[0].kwargs["params"]
    assert get.call_args_list[1].kwargs["params"]["page_token"] == "page2"
    post.assert_not_called()


@pytest.mark.parametrize("tokens", [[""], ["repeat", "repeat"]])
def test_bitable_broken_pagination_never_silently_returns_partial_report(monkeypatch, tokens):
    monkeypatch.setattr(bitable, "auth_headers", lambda: {})
    monkeypatch.setattr(bitable.requests, "get", Mock(side_effect=[
        response({"code": 0, "data": {"has_more": True, "page_token": token, "items": [row("a")]}})
        for token in tokens
    ]))
    with pytest.raises(RuntimeError, match="分页异常"):
        bitable.fetch_records("tbl-test", app_token="app-test")


def test_group_sender_uses_chat_id_and_all_http_is_mocked(monkeypatch):
    monkeypatch.setattr(feishu, "auth_headers", lambda: {})
    post = Mock(return_value=response({"code": 0, "data": {"message_id": "test"}}))
    monkeypatch.setattr(feishu.requests, "post", post)
    assert feishu.send_chat_text("oc-mock", "本地测试") == {"message_id": "test"}
    assert post.call_args.args[0].endswith("/im/v1/messages?receive_id_type=chat_id")
    assert post.call_args.kwargs["json"]["receive_id"] == "oc-mock"
    assert json.loads(post.call_args.kwargs["json"]["content"]) == {"text": "本地测试"}


@pytest.fixture
def sso_client(tmp_path, monkeypatch):
    auth = app_module.app.extensions["mhero_auth"]
    cfg = AuthConfig.from_env("nss", env={}, mode="on", issuer="http://127.0.0.1:18000",
                              client_id="mock", client_secret="mock", base_url="http://localhost",
                              session_secret="local-test-session-key-32-characters", public_get_paths=["/ds/brand/*"])
    monkeypatch.setattr(auth, "cfg", cfg)
    monkeypatch.setattr(app_module, "WEEKLY_REPORT_DIR", tmp_path)
    path = tmp_path / f"{WEEK}_weekly.xlsx"
    path.write_bytes(b"mock-excel")
    return app_module.app.test_client(), auth


@pytest.mark.parametrize("path", ["/weekly-reports", f"/weekly-reports/{WEEK}_weekly.xlsx"])
def test_weekly_report_list_and_download_require_existing_sso_login(sso_client, path):
    client, _ = sso_client
    for method in ("GET", "HEAD"):
        resp = client.open(path, method=method)
        assert resp.status_code == 302
        assert resp.location.startswith("/_auth/login?next=")
    resp = client.get(path, headers={"Accept": "application/json"})
    assert resp.status_code == 401


def login(client, auth, role="supervisor"):
    client.set_cookie(auth.cfg.cookie_name, auth._sign({
        "app": "nss", "sub": "local", "name": "local", "role": role,
        "exp": auth.cfg.clock() + 3600,
    }))


def test_logged_in_supervisor_can_list_download_but_not_escape_directory(sso_client, tmp_path):
    client, auth = sso_client
    login(client, auth)
    (tmp_path / "unrelated.xlsx").write_bytes(b"not-a-report")
    (tmp_path / f"{WEEK}_brief.txt").write_text("not-a-download")
    listing = client.get("/weekly-reports")
    assert listing.status_code == 200
    assert f"{WEEK}_weekly.xlsx" in listing.get_data(as_text=True)
    assert "unrelated.xlsx" not in listing.get_data(as_text=True)
    download = client.get(f"/weekly-reports/{WEEK}_weekly.xlsx")
    assert download.status_code == 200 and download.data == b"mock-excel"
    assert "attachment" in download.headers["Content-Disposition"]
    assert download.headers["Cache-Control"] == "no-store"
    for name in ("unrelated.xlsx", f"{WEEK}_brief.txt", "..%2Fsent_log.json"):
        assert client.get("/weekly-reports/" + name).status_code == 404
    (tmp_path / "2026-10-12_weekly.xlsx").symlink_to(tmp_path.parent / "outside.xlsx")
    assert client.get("/weekly-reports/2026-10-12_weekly.xlsx").status_code == 404


def test_logged_in_user_keeps_existing_role_restrictions(sso_client):
    client, auth = sso_client
    login(client, auth, "user")
    assert client.get("/weekly-reports").status_code == 403
    assert client.get(f"/weekly-reports/{WEEK}_weekly.xlsx").status_code == 403


def test_settings_api_rejects_conflicting_weekly_time_with_400():
    response = app_module.app.test_client().post("/api/settings", json={"weekly_report": {"hour": 12, "minute": 0}})
    assert response.status_code == 400
    assert "前后 5 分钟" in response.get_json()["error"]
